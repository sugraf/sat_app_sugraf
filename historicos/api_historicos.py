# historicos/api_historicos.py
"""Históricos SAT: lectura, búsqueda y API del histórico exportado desde Pirineos.

La carpeta "Historico SAT" de Drive contiene:
  - SATRealizados_<AÑO>.xlsx  -> un fichero por año
  - SATRealizadosActual.xlsx  -> partes del año en curso, se actualiza a menudo
  - SATRealizados.xlsx        -> histórico completo (solo se usa si no hay ficheros anuales)

Los partes se deduplican por N_Parte, prevaleciendo SATRealizadosActual.xlsx.
Los ficheros descargados se guardan en memoria mientras la instancia siga viva
y solo se vuelven a descargar si cambia su modifiedTime en Drive.

Lo usan la vista de Históricos (rutas /api/historicos/...) y la sección de
adjudicados (buscar_por_aviso).
"""
import io
import json
import os
import re
import threading
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import google_auth_httplib2
import httplib2
import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from google.oauth2 import service_account
from googleapiclient.discovery import build
from openpyxl import load_workbook

FOLDER_ID = "1o8Y2jrD86LBkxxmCnfH5tiacuFd1IJ51"  # Pirineos/Historico SAT
FICHERO_ACTUAL = 'SATRealizadosActual.xlsx'
FICHERO_COMPLETO = 'SATRealizados.xlsx'
PATRON_ANUAL = re.compile(r'^SATRealizados_(\d{4})\.xlsx$', re.IGNORECASE)
MIME_SHEETS = 'application/vnd.google-apps.spreadsheet'
MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_REGISTROS = 1000
MAX_POR_PAGINA = 200

FORMAS_JURIDICAS = {
    'SA', 'SL', 'SLU', 'SAU', 'SLL', 'SLNE', 'SC', 'SCP', 'SCOOP', 'SCL', 'CB', 'COOP',
}

router = APIRouter()

_lock = threading.Lock()
_cache_ficheros = {}  # file_id -> (modifiedTime, [registros])
_cache_indice = {}  # version (ids + modifiedTime) -> IndiceHistorico ya construido


def _credenciales():
    creds_dict = json.loads(os.environ.get("GOOGLE_CREDENTIALS_JSON"))
    return service_account.Credentials.from_service_account_info(
        creds_dict, scopes=["https://www.googleapis.com/auth/drive"]
    )


def _ficheros_a_leer(drive, folder_id):
    """Devuelve los ficheros del histórico en orden de prioridad creciente."""
    res = drive.files().list(
        q=f"'{folder_id}' in parents and trashed = false",
        fields="files(id, name, mimeType, modifiedTime)",
        pageSize=200,
        supportsAllDrives=True,
        includeItemsFromAllDrives=True,
    ).execute()
    por_nombre = {f['name']: f for f in res.get('files', [])}

    anuales = sorted(
        (f for nombre, f in por_nombre.items() if PATRON_ANUAL.match(nombre)),
        key=lambda f: f['name'],
    )
    ficheros = anuales or ([por_nombre[FICHERO_COMPLETO]] if FICHERO_COMPLETO in por_nombre else [])
    if FICHERO_ACTUAL in por_nombre:
        ficheros.append(por_nombre[FICHERO_ACTUAL])
    return ficheros


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d')
    return _texto(valor)[:10]


def _leer_registros(contenido):
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    filas = wb.worksheets[0].iter_rows(values_only=True)
    cabecera = [_texto(c) for c in next(filas, [])]
    registros = []
    for fila in filas:
        d = dict(zip(cabecera, fila))
        n_parte = _texto(d.get('N_Parte'))
        if not n_parte:
            continue
        registros.append({
            'n_parte': n_parte,
            'fecha_entrada': _fecha(d.get('Fecha_Entrada')),
            'fecha': _fecha(d.get('Fecha_Realizado')) or _fecha(d.get('Fecha_Entrada')),
            'cliente': _texto(d.get('Nombre')),
            'nombre_comercial': _texto(d.get('Nombre_Comercial')),
            'equipo': _texto(d.get('Equipo')),
            'marca': _texto(d.get('Marca')),
            'modelo': _texto(d.get('Modelo')),
            'n_serie': _texto(d.get('N_Serie')),
            'tecnico': _texto(d.get('Tecnico')),
            'anomalia': _texto(d.get('Anomalia')),
            'solucion': _texto(d.get('Solucion')),
        })
    wb.close()
    return registros


def _descargar(drive, creds, fichero):
    if fichero.get('mimeType') == MIME_SHEETS:
        peticion = drive.files().export_media(fileId=fichero['id'], mimeType=MIME_XLSX)
    else:
        peticion = drive.files().get_media(fileId=fichero['id'], supportsAllDrives=True)
    # httplib2 no es thread-safe: cada descarga usa su propia conexión.
    http = google_auth_httplib2.AuthorizedHttp(creds, http=httplib2.Http(timeout=60))
    return _leer_registros(peticion.execute(http=http))


def cargar_historico():
    """Devuelve (IndiceHistorico, nombres_de_ficheros) con todo el histórico deduplicado."""
    creds = _credenciales()
    drive = build("drive", "v3", credentials=creds, cache_discovery=False)
    ficheros = _ficheros_a_leer(drive, FOLDER_ID)
    if not ficheros:
        raise RuntimeError(
            "No se encuentran los Excel en la carpeta Historico SAT de Drive. "
            "Comprueba que está compartida con la cuenta de servicio de la app."
        )
    nombres = [f['name'] for f in ficheros]
    version = tuple((f['id'], f.get('modifiedTime')) for f in ficheros)

    with _lock:
        if _cache_indice.get('version') == version:
            return _cache_indice['indice'], nombres
        pendientes = [
            f for f in ficheros
            if _cache_ficheros.get(f['id'], (None,))[0] != f.get('modifiedTime')
        ]
    if pendientes:
        with ThreadPoolExecutor(max_workers=8) as pool:
            descargados = list(pool.map(lambda f: _descargar(drive, creds, f), pendientes))
        with _lock:
            for f, registros in zip(pendientes, descargados):
                _cache_ficheros[f['id']] = (f.get('modifiedTime'), registros)

    por_parte = {}
    with _lock:
        for f in ficheros:
            for r in _cache_ficheros[f['id']][1]:
                por_parte[r['n_parte']] = r
    indice = IndiceHistorico(list(por_parte.values()))
    with _lock:
        _cache_indice.update(version=version, indice=indice)
    return indice, nombres


# ---------------------------------------------------------------------------
# Normalización y búsqueda
# ---------------------------------------------------------------------------

def _sin_acentos(texto):
    texto = unicodedata.normalize('NFKD', _texto(texto))
    return ''.join(c for c in texto if not unicodedata.combining(c)).upper()


def _alfanum(texto):
    return re.sub(r'[^A-Z0-9]', '', _sin_acentos(texto))


def clave_cliente(nombre):
    """'GISTEL, S.L. INDUSTRIAS GRAFICAS' -> 'GISTEL INDUSTRIAS GRAFICAS'."""
    texto = _sin_acentos(nombre)
    # Une las siglas con puntos (S.L.U. -> SLU) antes de separar en palabras.
    texto = re.sub(r'\b((?:[A-Z]\.){1,3}[A-Z]?)\.?(?=\W|$)', lambda m: m.group(1).replace('.', ''), texto)
    palabras = re.findall(r'[A-Z0-9]+', texto)
    return ' '.join(p for p in palabras if p not in FORMAS_JURIDICAS)


def clave_serie(serie):
    valor = _texto(serie)
    if valor.endswith('.0'):
        valor = valor[:-2]
    valor = _alfanum(valor).lstrip('0')
    return valor if len(valor) >= 3 else ''


def clave_modelo(modelo):
    valor = _alfanum(modelo)
    return valor if len(valor) >= 3 else ''


class IndiceHistorico:
    """Registros del histórico con sus claves de búsqueda precalculadas."""

    def __init__(self, registros):
        self.registros = registros
        self.claves = [
            {
                'clientes': {clave_cliente(r['cliente']), clave_cliente(r['nombre_comercial'])} - {''},
                'serie': clave_serie(r['n_serie']),
                'modelo': clave_modelo(r['modelo']),
                # Textos normalizados para los filtros de la vista de Históricos.
                'txt_cliente': _sin_acentos(f"{r['cliente']} {r['nombre_comercial']}"),
                'txt_maquina': _sin_acentos(f"{r['equipo']} {r['marca']} {r['modelo']} {r['n_serie']}"),
                'alf_maquina': _alfanum(f"{r['modelo']} {r['n_serie']}"),
                'txt_libre': _sin_acentos(' '.join(
                    r[c] for c in ('n_parte', 'cliente', 'nombre_comercial', 'equipo', 'marca',
                                   'modelo', 'n_serie', 'tecnico', 'anomalia', 'solucion')
                )),
            }
            for r in registros
        ]
        # Un "modelo" que coincide con un tipo de equipo (GRAPADORA, RETRACTILADORA...)
        # es genérico y no sirve para buscar la misma máquina en otros clientes.
        self.modelos_genericos = {_alfanum(r['equipo']) for r in registros} - {''}
        self.modelos = {k['modelo'] for k in self.claves} - {''} - self.modelos_genericos


def opciones(indice):
    """Valores distintos para rellenar los desplegables de filtros."""
    def distintos(campo):
        conteo = {}
        for r in indice.registros:
            if r[campo]:
                conteo[r[campo]] = conteo.get(r[campo], 0) + 1
        return [{'valor': v, 'total': n} for v, n in sorted(conteo.items())]

    fechas = sorted(r['fecha'] for r in indice.registros if r['fecha'])
    return {
        'total': len(indice.registros),
        'clientes': distintos('cliente'),
        'tecnicos': distintos('tecnico'),
        'equipos': distintos('equipo'),
        'marcas': distintos('marca'),
        'fecha_min': fechas[0] if fechas else '',
        'fecha_max': fechas[-1] if fechas else '',
    }


def filtrar(indice, texto='', cliente='', maquina='', tecnico='', equipo='', marca='',
            desde='', hasta='', orden='desc'):
    """Filtro general de la vista de Históricos. Devuelve las posiciones que cumplen
    todos los filtros, ordenadas por fecha. Los textos ignoran acentos y mayúsculas;
    'texto' y 'cliente' exigen que aparezcan todas sus palabras."""
    palabras = _sin_acentos(texto).split()
    palabras_cliente = _sin_acentos(cliente).split()
    maquina_n = _sin_acentos(maquina)
    maquina_a = _alfanum(maquina)
    tecnico, equipo, marca = _texto(tecnico), _texto(equipo), _texto(marca)

    posiciones = []
    for i, (r, k) in enumerate(zip(indice.registros, indice.claves)):
        if desde and r['fecha'] < desde:
            continue
        if hasta and r['fecha'] > hasta:
            continue
        if tecnico and r['tecnico'] != tecnico:
            continue
        if equipo and r['equipo'] != equipo:
            continue
        if marca and r['marca'] != marca:
            continue
        if palabras_cliente and not all(p in k['txt_cliente'] for p in palabras_cliente):
            continue
        # "VG3-640" encuentra "VG3640" y viceversa.
        if maquina_n and maquina_n not in k['txt_maquina'] and not (maquina_a and maquina_a in k['alf_maquina']):
            continue
        if palabras and not all(p in k['txt_libre'] for p in palabras):
            continue
        posiciones.append(i)

    registros = indice.registros
    posiciones.sort(key=lambda i: (registros[i]['fecha'], registros[i]['n_parte']), reverse=(orden != 'asc'))
    return posiciones


def _contiene_cliente(claves_registro, clave):
    """Coincidencia laxa: uno de los nombres contiene al otro como palabras completas."""
    for otra in claves_registro:
        if len(otra) >= 6 and (f' {otra} ' in f' {clave} ' or f' {clave} ' in f' {otra} '):
            return True
    return False


def _marcas_compatibles(marca_a, marca_b):
    a, b = _alfanum(marca_a), _alfanum(marca_b)
    return not a or not b or a in b or b in a


def _modelo_desde_nombre(indice, maquina):
    """Busca el modelo más largo (con algún dígito) contenido en el nombre de la máquina."""
    texto = _alfanum(maquina)
    encontrados = [m for m in indice.modelos if len(m) >= 4 and re.search(r'\d', m) and m in texto]
    return max(encontrados, key=len) if encontrados else ''


def buscar_por_aviso(indice, cliente='', maquina='', marca='', modelo='', n_serie=''):
    registros, claves = indice.registros, indice.claves
    # "SIJALON / SUMINISTROS INDUSTRIALES JALON SL": nombre comercial / razón social.
    claves_cli = {clave_cliente(p) for p in [cliente, *_texto(cliente).split('/')]} - {''}
    serie = clave_serie(n_serie)

    # --- Histórico del cliente (todas sus máquinas) ---
    pos_cliente = []
    criterio_cliente = ''
    if claves_cli:
        pos_cliente = [i for i, k in enumerate(claves) if k['clientes'] & claves_cli]
        criterio_cliente = 'nombre exacto'
        if not pos_cliente:
            pos_cliente = [
                i for i, k in enumerate(claves)
                if any(_contiene_cliente(k['clientes'], c) for c in claves_cli)
            ]
            criterio_cliente = 'nombre aproximado'
    set_cliente = set(pos_cliente)

    # --- Histórico de la máquina (misma unidad + mismo modelo en otros clientes) ---
    pos_unidad = [i for i, k in enumerate(claves) if serie and k['serie'] == serie]
    set_unidad = set(pos_unidad)
    modelos = {clave_modelo(modelo)} | {claves[i]['modelo'] for i in pos_unidad}
    modelos -= {''} | indice.modelos_genericos

    def por_modelo(modelos):
        return [
            i for i, k in enumerate(claves)
            if i not in set_unidad and k['modelo'] in modelos
            and _marcas_compatibles(registros[i]['marca'], marca)
        ]

    pos_modelo = por_modelo(modelos) if modelos else []
    if not pos_unidad and not pos_modelo and maquina:
        # Sin resultados por serie/modelo: se deduce el modelo del nombre de la máquina.
        deducido = _modelo_desde_nombre(indice, maquina)
        if deducido:
            modelos = {deducido}
            pos_modelo = por_modelo(modelos)

    criterio_maquina = []
    if serie:
        criterio_maquina.append(f'nº serie {n_serie}')
    if modelos:
        criterio_maquina.append('modelo ' + ', '.join(sorted(modelos)))
    pos_maquina = pos_unidad + pos_modelo

    def anotar(posiciones):
        posiciones = sorted(
            posiciones,
            key=lambda i: (registros[i]['fecha'], registros[i]['n_parte']),
            reverse=True,
        )
        return [
            {**registros[i], 'misma_unidad': i in set_unidad, 'mismo_cliente': i in set_cliente}
            for i in posiciones[:MAX_REGISTROS]
        ], len(posiciones)

    reg_maquina, total_maquina = anotar(pos_maquina)
    reg_cliente, total_cliente = anotar(pos_cliente)
    return {
        'maquina': {
            'registros': reg_maquina,
            'total': total_maquina,
            'otros_clientes': sum(1 for i in pos_maquina if i not in set_cliente),
            'criterio': ' / '.join(criterio_maquina),
        },
        'cliente': {
            'registros': reg_cliente,
            'total': total_cliente,
            'criterio': criterio_cliente,
        },
    }


# ---------------------------------------------------------------------------
# API de la vista de Históricos
# ---------------------------------------------------------------------------

COLUMNAS_EXPORTAR = [
    ('n_parte', 'Nº Parte'), ('fecha_entrada', 'Fecha Entrada'), ('fecha', 'Fecha Realizado'),
    ('cliente', 'Cliente'), ('nombre_comercial', 'Nombre Comercial'), ('equipo', 'Equipo'),
    ('marca', 'Marca'), ('modelo', 'Modelo'), ('n_serie', 'Nº Serie'), ('tecnico', 'Técnico'),
    ('anomalia', 'Anomalía'), ('solucion', 'Solución'),
]


def _cargar_y_filtrar(texto, cliente, maquina, tecnico, equipo, marca, desde, hasta, orden):
    indice, _ = cargar_historico()
    posiciones = filtrar(
        indice, texto=texto, cliente=cliente, maquina=maquina, tecnico=tecnico,
        equipo=equipo, marca=marca, desde=desde, hasta=hasta, orden=orden,
    )
    return indice, posiciones


@router.get("/opciones")
def get_opciones():
    try:
        indice, fuentes = cargar_historico()
        resultado = opciones(indice)
        resultado['fuentes'] = fuentes
        return resultado
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo leer el histórico de SAT: {e}")


@router.get("/buscar")
def get_buscar(texto: str = '', cliente: str = '', maquina: str = '', tecnico: str = '', equipo: str = '',
               marca: str = '', desde: str = '', hasta: str = '', orden: str = 'desc',
               pagina: int = 1, por_pagina: int = 50):
    try:
        indice, posiciones = _cargar_y_filtrar(texto, cliente, maquina, tecnico, equipo, marca, desde, hasta, orden)
        por_pagina = max(1, min(por_pagina, MAX_POR_PAGINA))
        pagina = max(1, pagina)
        inicio = (pagina - 1) * por_pagina
        registros = [indice.registros[i] for i in posiciones[inicio:inicio + por_pagina]]
        return {
            'total': len(posiciones),
            'clientes': len({indice.registros[i]['cliente'] for i in posiciones}),
            'pagina': pagina,
            'por_pagina': por_pagina,
            'registros': registros,
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo buscar en el histórico de SAT: {e}")


@router.get("/exportar")
def get_exportar(texto: str = '', cliente: str = '', maquina: str = '', tecnico: str = '', equipo: str = '',
                 marca: str = '', desde: str = '', hasta: str = '', orden: str = 'desc'):
    try:
        indice, posiciones = _cargar_y_filtrar(texto, cliente, maquina, tecnico, equipo, marca, desde, hasta, orden)
        df = pd.DataFrame(
            [[indice.registros[i][campo] for campo, _ in COLUMNAS_EXPORTAR] for i in posiciones],
            columns=[titulo for _, titulo in COLUMNAS_EXPORTAR],
        )
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Historico SAT')
        output.seek(0)
        nombre = f"Historico_SAT_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return StreamingResponse(
            output,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{nombre}"'},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo exportar el histórico de SAT: {e}")
